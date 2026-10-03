"""Excel exports go through the same pipeline as CSV (issue #36): same rules, reasons and idempotency."""

import csv
import zipfile
from datetime import date, datetime, time
from io import BytesIO, StringIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.catalog.models import KNOWN_ROBOTS, Episode, ImportBatch, Robot
from tests.test_import_service import HEADER, ROW_1, SEED_FILE

IMPORTS = "/api/imports/"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def robots():
    for robot_id in KNOWN_ROBOTS:
        Robot.objects.create(id=robot_id, kind=Robot.kind_from_id(robot_id))


def workbook(rows) -> bytes:
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(list(row))
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def upload(client, content: bytes, name="export.xlsx", content_type=XLSX):
    return client.post(IMPORTS, {"file": SimpleUploadedFile(name, content, content_type)}, format="multipart")


COUNTS = ("total_rows", "created_count", "updated_count", "unchanged_count", "skipped_count", "fixed_count")


def test_the_seed_export_saved_as_xlsx_gives_the_same_report_as_the_csv(api_as):
    operator = api_as("operator")
    csv_report = upload(operator, SEED_FILE.read_bytes(), name="episodes.csv", content_type="text/csv").json()
    csv_issues = [
        (i["row_number"], i["reason_code"])
        for i in operator.get(f"{IMPORTS}{csv_report['id']}/issues/?page_size=100").json()["results"]
    ]
    Episode.objects.all().delete()

    with SEED_FILE.open(encoding="utf-8-sig", newline="") as export:
        rows = list(csv.reader(export))
    xlsx_report = upload(operator, workbook(rows), name="episodes.xlsx").json()
    xlsx_issues = [
        (i["row_number"], i["reason_code"])
        for i in operator.get(f"{IMPORTS}{xlsx_report['id']}/issues/?page_size=100").json()["results"]
    ]

    assert xlsx_report["status"] == "completed"
    assert {key: xlsx_report[key] for key in COUNTS} == {key: csv_report[key] for key in COUNTS}
    # Same reasons on the same rows. One exception: CSV can tell a line is short (too few columns), but a
    # sheet has no short rows, only empty cells, so it reports the first missing value. Skipped either way.
    assert [row for row, _ in xlsx_issues] == [row for row, _ in csv_issues]
    for (row, xlsx_code), (_, csv_code) in zip(xlsx_issues, csv_issues, strict=True):
        assert xlsx_code == csv_code or (csv_code == "malformed_row" and xlsx_code.startswith("missing_")), (
            row
        )


def test_real_excel_dates_and_numbers_are_understood(api_as):
    content = workbook(
        [
            HEADER.split(","),
            ["EP-00001", "arm-01", "pick cup", datetime(2026, 8, 1, 10, 0), 30.0, "Aline", "good"],
        ]
    )

    report = upload(api_as("operator"), content).json()

    assert (report["created_count"], report["skipped_count"], report["fixed_count"]) == (1, 0, 0)
    episode = Episode.objects.get()
    assert episode.duration_seconds == 30
    assert episode.recorded_at.isoformat().startswith("2026-08-01T10:00")


def test_the_same_workbook_uploaded_again_is_flagged_and_creates_nothing(api_as):
    operator = api_as("operator")
    content = workbook([HEADER.split(","), ROW_1.split(",")])
    upload(operator, content)

    report = upload(operator, content).json()

    assert report["previously_imported"] is True
    assert (report["created_count"], report["unchanged_count"]) == (0, 1)


def test_row_numbers_in_the_report_are_the_sheet_rows(api_as):
    operator = api_as("operator")
    content = workbook(
        [HEADER.split(","), ROW_1.split(","), ["EP-00009", "arm-99", "x", "2026-08-01", 5, "A", "good"]]
    )

    report = upload(operator, content).json()

    issues = operator.get(f"{IMPORTS}{report['id']}/issues/").json()["results"]
    assert [(issue["row_number"], issue["reason_code"]) for issue in issues] == [(3, "unknown_robot")]


def test_empty_formatted_rows_after_the_data_are_ignored(api_as):
    book = Workbook()
    sheet = book.active
    sheet.append(HEADER.split(","))
    sheet.append(ROW_1.split(","))
    sheet.cell(row=40, column=1).font = Font(bold=True)  # formatting only: Excel still counts the row
    out = BytesIO()
    book.save(out)

    report = upload(api_as("operator"), out.getvalue()).json()

    assert (report["total_rows"], report["skipped_count"]) == (1, 0)


def test_a_file_that_is_not_a_workbook_fails_cleanly(api_as):
    response = upload(api_as("operator"), b"not a zip at all", name="export.xlsx")

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "invalid_file"
    assert "Excel workbook" in response.json()["error"]["message"]
    assert ImportBatch.objects.get().status == "failed"


def test_a_workbook_that_unpacks_too_large_is_refused(api_as, settings):
    settings.IMPORT_MAX_XLSX_UNPACKED_BYTES = 1000  # any real workbook is bigger than this once unzipped

    response = upload(api_as("operator"), workbook([HEADER.split(","), ROW_1.split(",")]))

    assert response.status_code == 400
    assert "too large" in response.json()["error"]["message"]
    assert Episode.objects.count() == 0


@pytest.mark.parametrize("name", ["export.xls", "export.ods", "export.pdf"])
def test_other_file_types_are_refused_with_what_is_accepted(api_as, name):
    response = upload(api_as("operator"), b"anything", name=name, content_type="application/octet-stream")

    assert response.status_code == 400
    assert response.json()["error"]["details"]["file"] == ["Upload the export as a .csv or .xlsx file."]


def test_the_command_imports_a_workbook(make_user, tmp_path):
    make_user(email="ops@example.com", role="operator")
    path = tmp_path / "episodes.xlsx"
    path.write_bytes(workbook([HEADER.split(","), ROW_1.split(",")]))
    out = StringIO()

    call_command("import_episodes", str(path), "--as", "ops@example.com", stdout=out)
    call_command("import_episodes", str(path), "--as", "ops@example.com", "--once", stdout=out)

    assert "1 created" in out.getvalue()
    assert "already imported" in out.getvalue()
    assert ImportBatch.objects.count() == 1


def test_a_zip_that_is_not_a_workbook_fails_cleanly(api_as):
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w") as zipped:
        zipped.writestr("notes.txt", "hello")

    response = upload(api_as("operator"), archive.getvalue())

    assert response.status_code == 400
    assert "Excel workbook" in response.json()["error"]["message"]


def test_date_only_cells_and_text_with_line_breaks_are_read_as_the_csv_would_be():
    from apps.catalog.spreadsheet import _text

    assert _text(date(2026, 8, 1)) == "2026-08-01"
    assert _text(time(10, 30)) == "10:30:00"
    assert _text("Aline\nUwase") == "Aline Uwase"
    assert _text(30.5) == "30.5"
