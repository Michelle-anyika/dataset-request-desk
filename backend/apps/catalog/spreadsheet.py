"""Read an export as text lines, whether it is CSV or an Excel workbook (issue #36).

A workbook's first sheet is turned into CSV lines, one per sheet row, and handed to the same import pipeline:
the same normalisation, report reasons and idempotency, and row N of the sheet is line N of the report.
"""

import csv
import io
import zipfile
from collections.abc import Iterator
from datetime import date, time
from pathlib import Path
from typing import BinaryIO

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.catalog.import_service import ImportFailed

EXCEL = ".xlsx"
NOT_A_WORKBOOK = "The file is not a valid Excel workbook. Save it as .xlsx, or export it as CSV."


def export_lines(file: BinaryIO, name: str) -> Iterator[str]:
    """The export's text lines: CSV as uploaded, or a workbook's first sheet converted to CSV."""
    if Path(name).suffix.lower() == EXCEL:
        return xlsx_lines(file)
    # utf-8-sig drops a byte-order mark if the export has one; newline="" lets the csv module see CRLF.
    return io.TextIOWrapper(file, encoding="utf-8-sig", newline="")


def xlsx_lines(file: BinaryIO) -> Iterator[str]:
    _refuse_zip_bombs(file)
    try:
        book = load_workbook(file, read_only=True, data_only=True)  # values, not formulas
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise ImportFailed(NOT_A_WORKBOOK) from exc
    try:
        for values in book.worksheets[0].iter_rows(values_only=True):
            yield _csv_line(values)
    finally:
        book.close()


def _refuse_zip_bombs(file: BinaryIO) -> None:
    """A workbook is a zip, so a small upload can unpack to gigabytes: check the sizes before opening it."""
    try:
        with zipfile.ZipFile(file) as archive:
            unpacked = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ImportFailed(NOT_A_WORKBOOK) from exc
    finally:
        file.seek(0)
    if unpacked > settings.IMPORT_MAX_XLSX_UNPACKED_BYTES:
        limit_mb = settings.IMPORT_MAX_XLSX_UNPACKED_BYTES / (1024 * 1024)
        raise ImportFailed(f"The workbook is too large: it unpacks to more than {limit_mb:g} MB.")


def _csv_line(values) -> str:
    if all(value is None for value in values):
        # An empty row is an empty line, as in a CSV export: the reader skips it without counting it, so
        # formatted but empty rows (Excel keeps those) never show up in the report.
        return "\n"
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerow([_text(value) for value in values])
    return out.getvalue()


def _text(value) -> str:
    """A cell as the CSV export would have written it, so the existing parsing rules apply unchanged."""
    if value is None:
        return ""
    if isinstance(value, date | time):  # dates, times and datetimes: ISO 8601, which the parser reads first
        return value.isoformat()
    # Whole numbers arrive as int (30, not 30.0), so str() writes them as the CSV export would.
    # A line break inside a cell would shift every later line number in the report; episodes have no
    # multi-line fields, so it is read as a space.
    return " ".join(str(value).splitlines())
